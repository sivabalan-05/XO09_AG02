from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

OUT = '/private/tmp/Kabir_Sen_Full_Match_CV.docx'

def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd'); shd.set(qn('w:fill'), fill); tc_pr.append(shd)

def set_cell_border(cell, color='D9D9D9'):
    tc_pr = cell._tc.get_or_add_tcPr(); borders = OxmlElement('w:tcBorders')
    for edge in ('top', 'left', 'bottom', 'right'):
        tag = OxmlElement(f'w:{edge}'); tag.set(qn('w:val'), 'single'); tag.set(qn('w:sz'), '4'); tag.set(qn('w:color'), color); borders.append(tag)
    tc_pr.append(borders)

doc = Document()
sec = doc.sections[0]
sec.top_margin = Inches(.58); sec.bottom_margin = Inches(.55); sec.left_margin = Inches(.72); sec.right_margin = Inches(.72)

styles = doc.styles
styles['Normal'].font.name = 'Aptos'; styles['Normal'].font.size = Pt(10); styles['Normal'].font.color.rgb = RGBColor(36, 48, 42)
for name in ('Title', 'Heading 1'):
    styles[name].font.name = 'Aptos Display'; styles[name].font.color.rgb = RGBColor(0, 0, 0)

title = doc.add_paragraph(); title.style = 'Title'; title.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = title.add_run('Kabir Sen'); r.bold = True; r.font.size = Pt(26)
sub = doc.add_paragraph(); sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = sub.add_run('Platform Engineer  |  Bengaluru, India  |  kabir.sen@example.com  |  github.com/kabir-sen'); r.font.size = Pt(9.5); r.font.color.rgb = RGBColor(74, 91, 81)

def heading(text):
    p = doc.add_paragraph(); p.style = 'Heading 1'; p.paragraph_format.space_before = Pt(12); p.paragraph_format.space_after = Pt(4)
    run = p.add_run(text.upper()); run.font.size = Pt(11); run.bold = True; run.font.color.rgb = RGBColor(32, 83, 60)
    return p

def bullet(text):
    p = doc.add_paragraph(style='List Bullet'); p.paragraph_format.space_after = Pt(2); p.paragraph_format.left_indent = Inches(.18); p.paragraph_format.first_line_indent = Inches(-.12)
    p.add_run(text)

heading('Profile')
p = doc.add_paragraph('Platform engineer with four years of hands-on backend, cloud, reliability, and infrastructure-as-code experience. Led production multi-region failover for event-driven services serving 18,000 requests per second, with documented recovery objectives and incident ownership.')
p.paragraph_format.space_after = Pt(2)

heading('Experience')
p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(3)
r = p.add_run('Platform Engineer  |  Northstar Commerce  |  2022 – Present'); r.bold = True; r.font.size = Pt(10.5)
bullet('Led the design and quarterly failover exercises for active-active Java services on AWS EKS across Mumbai and Singapore, sustaining 18,000 requests per second and restoring critical traffic in 11 minutes during a regional simulation.')
bullet('Built and operated Kafka-backed order and inventory services; introduced idempotency and consumer lag controls that reduced duplicate-event incidents by 87%.')
bullet('Wrote Terraform modules for EKS, VPC networking, IAM, and regional routing; enabled reproducible environments for six product teams.')
bullet('Introduced OpenTelemetry tracing, Prometheus service-level objectives, Grafana dashboards, and paging alerts for 24 services; reduced mean time to diagnose production incidents from 31 to 12 minutes.')
bullet('Owned incident reviews, wrote operational runbooks, and partnered with security, support, and product teams during releases.')

p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(5); p.paragraph_format.space_after = Pt(3)
r = p.add_run('Software Engineer  |  Atlas Systems  |  2021 – 2022'); r.bold = True; r.font.size = Pt(10.5)
bullet('Developed Java REST APIs and asynchronous workers for a logistics platform, deployed through Docker and Kubernetes on AWS.')
bullet('Added structured logging, automated integration tests, and release checks for customer-facing services.')

heading('Selected Project')
p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(2)
r = p.add_run('Regional Resilience Drill Toolkit'); r.bold = True
p.add_run('  |  Internal platform project')
bullet('Created a reusable Kubernetes and Terraform test harness that validates DNS failover, queue replay, health checks, and recovery-time objectives before production disaster-recovery drills.')

heading('Skills')
table = doc.add_table(rows=2, cols=4)
table.autofit = False
skills = [('Backend', 'Java, REST APIs, microservices'), ('Distributed systems', 'Kafka, idempotency, event-driven services'), ('Cloud and platform', 'AWS, EKS, Docker, multi-region failover'), ('Reliability', 'OpenTelemetry, Prometheus, Grafana, SLOs'), ('Infrastructure as code', 'Terraform, IAM, VPC networking'), ('Collaboration', 'Runbooks, incident reviews, stakeholder delivery'), ('Databases', 'PostgreSQL, Redis'), ('Delivery', 'GitHub Actions, automated tests')]
for idx, (label, value) in enumerate(skills):
    cell = table.cell(idx // 4, idx % 4); set_cell_border(cell); shade(cell, 'F4F7F4')
    cell.width = Inches(1.72)
    p = cell.paragraphs[0]; p.paragraph_format.space_after = Pt(1)
    run = p.add_run(label); run.bold = True; run.font.size = Pt(8); run.font.color.rgb = RGBColor(32, 83, 60)
    p = cell.add_paragraph(value); p.paragraph_format.space_after = Pt(2); p.runs[0].font.size = Pt(8)

heading('Education')
p = doc.add_paragraph('B.Tech, Computer Science  |  2021'); p.paragraph_format.space_after = Pt(0)

footer = sec.footer.paragraphs[0]; footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = footer.add_run('Synthetic hackathon demonstration CV  •  Evidence is intentionally written for the PS02 screening scenario'); r.font.size = Pt(7.5); r.font.color.rgb = RGBColor(110, 120, 113)

doc.save(OUT)
print(OUT)
